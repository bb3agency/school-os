"""UDISE+ pre-check and student list carry PEN and APAAR ID (ADR-0037; FR-EXP-005, PRV-020,
invariant 4). The typed APAAR column is written unmasked; every other cell keeps the Aadhaar
mask. Synthetic data only: the Verhoeff-valid number is built inside the test process."""

from __future__ import annotations

import csv
import io
import sys
import uuid
from typing import Any

import pytest
from openpyxl import load_workbook
from sqlalchemy import Engine, text

from app.core.db import tenant_session
from app.core.redaction import verhoeff_check_digit, verhoeff_valid
from app.exports.config import load_config
from app.exports.tables import Table, TypedDigits12, safe_cell, typed_digits12, write_csv
from app.students.schemas import ValueIn

EX = sys.modules["sos_test_exports_objects"]


def verhoeff_number(body: str = "67812345678") -> str:
    # FR-STU-018: one APAAR ID per live student, so each test uses its own number.
    number = body + verhoeff_check_digit(body)
    assert verhoeff_valid(number)
    return number


# --- pure ----------------------------------------------------------------------------------------


def test_FR_EXP_005_udise_layout_has_pen_and_apaar_columns() -> None:
    layout = load_config().profiles["udise-plus"]
    assert layout.layout_version >= 2  # 3 since ADR-0039 (2026-27 portal order)
    assert layout.fields[-2:] == ("udise_pen", "apaar_id")
    assert "apaar_id" not in load_config().profiles["cisce-registration-2026"].fields


def test_PRV_020_only_the_typed_apaar_value_skips_the_mask() -> None:
    number = verhoeff_number()
    typed = typed_digits12(number)
    assert isinstance(typed, TypedDigits12)
    assert safe_cell(typed) == number
    # The same digits as plain text (any other column or free text) are masked.
    assert safe_cell(number) == f"XXXX XXXX {number[-4:]}"
    assert safe_cell(f"APAAR {number}") == f"APAAR XXXX XXXX {number[-4:]}"
    # Anything that is not exactly 12 digits never becomes typed.
    for raw in (f"APAAR {number}", f"{number[:4]} {number[4:8]} {number[8:]}", number + "1"):
        assert not isinstance(typed_digits12(raw), TypedDigits12)
        assert safe_cell(typed_digits12(raw)) == safe_cell(raw)
    with pytest.raises(ValueError, match="12 ASCII digits"):
        TypedDigits12(f"x{number}")


def test_PRV_020_csv_keeps_typed_apaar_and_masks_the_rest() -> None:
    number = verhoeff_number()
    table = Table(
        name="Ready",
        header=("APAAR ID", "Notes"),
        rows=((typed_digits12(number), f"seen {number}"),),
    )
    rows = list(csv.reader(io.StringIO(write_csv(table, watermark="W").decode("utf-8-sig"))))
    assert rows[2] == [number, f"seen XXXX XXXX {number[-4:]}"]


# --- end to end ----------------------------------------------------------------------------------


@pytest.fixture
def section(school: Any) -> str:
    from app.tenancy import service as tenancy
    from app.tenancy.schemas import SectionCreate

    with tenant_session(school.tenant_id, school.people["owner"].user_id) as db:
        sec = tenancy.create_section(
            db,
            SectionCreate(
                academic_year_id=school.ids["year"],
                class_id=school.ids["class_ix"],
                name=f"A{uuid.uuid4().hex[:6].upper()}",
            ),
        )
    key = f"section_{sec.id}"
    school.ids[key] = sec.id
    return key


def _plant_school_name(admin: Engine, school: Any, name: str) -> None:
    with admin.begin() as c:
        c.execute(
            text("UPDATE core.tenants SET name = :n WHERE id = :t"),
            {"n": name, "t": school.tenant_id},
        )


@pytest.mark.db
def test_FR_EXP_005_udise_precheck_ready_sheet_has_pen_and_unmasked_apaar(
    school: Any, section: str, admin_engine: Engine
) -> None:
    number = verhoeff_number()
    with_ids = EX.student(
        school,
        section_key=section,
        extra=[
            ValueIn(attribute_key="apaar_id", source="udise_plus", value=number),
            ValueIn(attribute_key="udise_pen", source="udise_plus", value="41345678903"),
        ],
    )
    EX.student(school, section_key=section)  # no IDs: empty cells
    admin = school.people["office_admin"]
    out = EX.request_precheck(
        school, admin, "office_admin", profile_key="udise-plus", section_keys=(section,)
    )
    # Free text in the same file (the school name) still masks the very same digits.
    _plant_school_name(admin_engine, school, f"Synthetic School {number}")
    try:
        assert EX.run(school, out.id) == "ready"
    finally:
        _plant_school_name(admin_engine, school, "Synthetic Model School")
    wb = load_workbook(io.BytesIO(EX.stored(out.id, "xlsx")))
    ready = [[c.value for c in row] for row in wb.worksheets[2].iter_rows()]
    header = ready[0]
    pen, apaar = header.index("UDISE+ PEN"), header.index("APAAR ID")
    assert apaar == pen + 1
    rows = {r[header.index("Admission number")]: r for r in ready[1:]}
    assert {(r[pen] or "", r[apaar] or "") for r in rows.values()} == {
        ("41345678903", number),
        ("", ""),
    }
    summary = [[c.value for c in row] for row in wb.worksheets[0].iter_rows()]
    flat = " ".join(str(v) for row in summary for v in row if v is not None)
    assert f"XXXX XXXX {number[-4:]}" in flat
    assert number not in flat
    assert number not in EX.FAKE.pages[-1], "the PDF has no APAAR column and masks free text"
    assert with_ids is not None


@pytest.mark.db
def test_FR_EXP_005_student_list_apaar_column_is_typed(school: Any, section: str) -> None:
    number = verhoeff_number("67812345679")
    EX.student(
        school,
        section_key=section,
        extra=[ValueIn(attribute_key="apaar_id", source="parent_form", value=number)],
    )
    out = EX.request_list(
        school,
        school.people["office_admin"],
        "office_admin",
        columns=("admission_no", "apaar_id"),
        section_keys=(section,),
    )
    assert EX.run(school, out.id) == "ready"
    rows = list(csv.reader(io.StringIO(EX.stored(out.id, "csv").decode("utf-8-sig"))))
    assert rows[1] == ["Admission number", "APAAR ID"]
    assert rows[2][1] == number
