"""Staged sheet of spreadsheet imports: view, edit, download (US-401 AC5/AC6; FR-IMP-008,
FR-IMP-009; invariants 4, 6, 7). Synthetic data only; Aadhaar-like numbers are built here.

Authorization for every role is in tests/security/test_authz_matrix.py and other-school ids in
tests/security/test_bola.py; these tests cover the behaviour.
"""

from __future__ import annotations

import csv
import io
import sys
import uuid
from typing import Any

import pytest
from openpyxl import load_workbook
from sqlalchemy import Engine, text

pytestmark = pytest.mark.db
S = sys.modules["sos_test_imports_support"]

TELUGU = "సింథెటిక విద్యార్థి"


def _sheet(api: Any, who: Any, batch_id: uuid.UUID, **params: Any) -> Any:
    res = api.call(who, "GET", f"/api/v1/imports/{batch_id}/sheet", params=params)
    assert res.status_code == 200, res.text
    return res


def _edit(
    api: Any, who: Any, batch_id: uuid.UUID, row_no: int, cells: list[dict[str, Any]], *, etag: str
) -> Any:
    return api.call(
        who,
        "PATCH",
        f"/api/v1/imports/{batch_id}/sheet/rows/{row_no}",
        json={"cells": cells},
        headers={"If-Match": etag},
    )


def _audit(admin: Engine, batch_id: uuid.UUID, action: str) -> list[dict[str, Any]]:
    with admin.connect() as c:
        rows = c.execute(
            text(
                "SELECT summary, actor_id FROM audit.events WHERE resource_id = :r "
                "AND action = :a ORDER BY seq"
            ),
            {"r": batch_id, "a": action},
        )
        return [dict(r._mapping) for r in rows]


def _batch_with_problem(admin: Engine, world: Any) -> tuple[uuid.UUID, list[str]]:
    rows, numbers = S.class_list(2)
    rows[0].append("Religion")  # a restricted (C3) field
    rows[1].append("Synthetic faith")
    rows[2].append("Synthetic faith")
    rows.append(["", "Synthetica Missing Number", "", "14/03/2012", "M", "IX", "A", ""])
    batch_id = S.start(admin, world.a, S.xlsx_bytes(rows))
    assert S.batch(admin, batch_id)["status"] == "validated"
    return batch_id, numbers


def test_FR_IMP_008_sheet_shows_every_column_and_row_with_checks(
    world: Any, api: Any, admin_engine: Engine
) -> None:
    admin = world.person("office_admin")
    batch_id, numbers = _batch_with_problem(admin_engine, world)
    res = _sheet(api, admin, batch_id)
    body = res.json()
    assert res.headers["ETag"] == f'W/"{body["version"]}"'
    assert body["editable"] is True
    assert body["read_only_reason"] is None
    assert body["header_row"] == 1
    assert body["total_rows"] == 3
    assert [c["letter"] for c in body["columns"]][:3] == ["A", "B", "C"]
    assert body["columns"][0]["header"] == "Adm No"
    assert body["columns"][0]["target"] == "admission_no"
    religion = body["columns"][7]
    assert religion["target"] == "religion"
    assert religion["restricted"] is True
    assert religion["editable"] is False
    first, _, third = body["data"]
    assert first["row_no"] == 2
    assert first["cells"][0]["value"] == numbers[0]
    assert first["status"] == "valid"
    # Restricted values are never sent (docs/05 §5.2): only that the cell is restricted.
    assert first["cells"][7] == {
        "value": None,
        "edited": False,
        "restricted": True,
        "formula": False,
    }
    assert "Synthetic faith" not in res.text
    assert third["status"] == "error"
    assert third["errors"][0]["code"] == "missing"
    page = _sheet(api, admin, batch_id, limit=2).json()
    assert [r["row_no"] for r in page["data"]] == [2, 3]
    rest = _sheet(api, admin, batch_id, limit=2, cursor=page["next_cursor"]).json()
    assert [r["row_no"] for r in rest["data"]] == [4]
    assert rest["offset"] == 2
    assert rest["next_cursor"] is None


def test_FR_IMP_008_edit_rechecks_the_row_and_commit_adds_the_edited_values(
    world: Any, api: Any, admin_engine: Engine
) -> None:
    admin = world.person("office_admin")
    batch_id, _ = _batch_with_problem(admin_engine, world)
    etag = _sheet(api, admin, batch_id).headers["ETag"]
    fixed_number = S.adm("EDT")
    res = _edit(
        api,
        admin,
        batch_id,
        4,
        [{"column": 0, "value": fixed_number}, {"column": 1, "value": f"  {TELUGU}  "}],
        etag=etag,
    )
    assert res.status_code == 200, res.text
    out = res.json()
    assert res.headers["ETag"] != etag
    assert out["row"]["status"] == "valid"
    assert out["row"]["errors"] == []
    assert out["row"]["cells"][0] == {
        "value": fixed_number,
        "edited": True,
        "restricted": False,
        "formula": False,
    }
    assert out["row"]["cells"][1]["value"] == TELUGU  # NFC, trimmed
    assert out["error_count"] == 0
    batch = S.batch(admin_engine, batch_id)
    assert batch["error_count"] == 0
    assert batch["stats"]["valid"] == 3
    # The history is ciphertext only (docs/05 §5.2.1): no plaintext value in the table.
    with admin_engine.connect() as c:
        edits = c.execute(
            text(
                "SELECT row_no, column_index, old_value_ciphertext, new_value_ciphertext, "
                "key_version, edited_by FROM sis.import_cell_edits WHERE batch_id = :b "
                "ORDER BY column_index"
            ),
            {"b": batch_id},
        ).all()
    assert [(e.row_no, e.column_index) for e in edits] == [(4, 0), (4, 1)]
    assert edits[0].old_value_ciphertext is None  # the cell was empty
    assert edits[1].old_value_ciphertext is not None
    for e in edits:
        assert e.edited_by == admin.user_id
        assert e.key_version >= 1
        assert fixed_number.encode() not in bytes(e.new_value_ciphertext)
        assert TELUGU.encode() not in bytes(e.new_value_ciphertext)
    # Audit per cell, field names only (invariant 7; never values).
    events = _audit(admin_engine, batch_id, "import.cell_edited")
    assert [(e["summary"]["row_no"], e["summary"]["field"]) for e in events] == [
        (4, "admission_no"),
        (4, "full_name"),
    ]
    assert fixed_number not in repr(events)
    assert TELUGU not in repr(events)
    # The sheet shows the edits as edited; the raw file is unchanged.
    again = _sheet(api, admin, batch_id).json()
    assert again["edited_cells"] == 2
    assert again["data"][2]["cells"][0]["edited"] is True
    # Commit re-reads the file with the edits applied: the edited row becomes a student.
    assert S.commit(world.a, batch_id) == "committed", S.batch(admin_engine, batch_id)
    student = S.student_by_adm(admin_engine, world.a.tenant_id, fixed_number)
    assert student is not None
    with admin_engine.connect() as c:
        names: list[str | None] = list(
            c.execute(
                text(
                    "SELECT value_text FROM sis.attribute_values WHERE student_id = :s "
                    "AND attribute_key = 'full_name'"
                ),
                {"s": student},
            ).scalars()
        )
    assert names == [TELUGU]


def test_FR_IMP_008_duplicate_fixed_in_one_row_clears_the_other(
    world: Any, api: Any, admin_engine: Engine
) -> None:
    admin = world.person("office_admin")
    rows, numbers = S.class_list(2)
    rows[2][0] = numbers[0]  # both rows carry the same admission number
    batch_id = S.start(admin_engine, world.a, S.xlsx_bytes(rows))
    sheet = _sheet(api, admin, batch_id).json()
    assert [r["status"] for r in sheet["data"]] == ["error", "error"]
    res = _edit(
        api,
        admin,
        batch_id,
        3,
        [{"column": 0, "value": S.adm("DUP")}],
        etag=f'W/"{sheet["version"]}"',
    )
    assert res.status_code == 200, res.text
    assert res.json()["row"]["status"] == "valid"
    assert res.json()["changed_rows"] == [2]
    assert [r["status"] for r in _sheet(api, admin, batch_id).json()["data"]] == [
        "valid",
        "valid",
    ]


def test_FR_IMP_008_edits_are_refused_safely(world: Any, api: Any, admin_engine: Engine) -> None:
    admin = world.person("office_admin")
    batch_id, _ = _batch_with_problem(admin_engine, world)
    etag = _sheet(api, admin, batch_id).headers["ETag"]
    before = S.count(
        admin_engine, "SELECT count(*) FROM sis.import_cell_edits WHERE batch_id = :b", b=batch_id
    )
    aadhaar = S.valid_aadhaar()
    res = _edit(api, admin, batch_id, 2, [{"column": 2, "value": aadhaar}], etag=etag)
    assert res.status_code == 422
    assert res.json()["errors"] == [
        {
            "field": "cells.0.value",
            "code": "aadhaar_full_number_rejected",
            "message_key": "errors.aadhaar_last4_only",
        }
    ]
    assert aadhaar not in res.text
    spaced = f"{aadhaar[:4]} {aadhaar[4:8]} {aadhaar[8:]}"
    assert (
        _edit(api, admin, batch_id, 2, [{"column": 2, "value": spaced}], etag=etag).status_code
        == 422
    )
    bad = _edit(api, admin, batch_id, 2, [{"column": 2, "value": "Line\nbreak"}], etag=etag)
    assert bad.json()["errors"][0]["code"] == "control_characters"
    restricted = _edit(api, admin, batch_id, 2, [{"column": 7, "value": "Other"}], etag=etag)
    assert restricted.status_code == 422
    assert restricted.json()["errors"][0]["code"] == "column_restricted"
    unknown = _edit(api, admin, batch_id, 2, [{"column": 40, "value": "x"}], etag=etag)
    assert unknown.json()["errors"][0]["code"] == "unknown_column"
    too_long = _edit(api, admin, batch_id, 2, [{"column": 2, "value": "x" * 1001}], etag=etag)
    assert too_long.status_code == 422
    missing_row = _edit(api, admin, batch_id, 99, [{"column": 2, "value": "x"}], etag=etag)
    assert missing_row.status_code == 404
    no_etag = api.call(
        admin,
        "PATCH",
        f"/api/v1/imports/{batch_id}/sheet/rows/2",
        json={"cells": [{"column": 2, "value": "x"}]},
    )
    assert no_etag.status_code == 400
    assert (
        S.count(
            admin_engine,
            "SELECT count(*) FROM sis.import_cell_edits WHERE batch_id = :b",
            b=batch_id,
        )
        == before
    )
    # Optimistic concurrency: the second edit with the same ETag answers 412.
    ok = _edit(api, admin, batch_id, 2, [{"column": 2, "value": "Synthetic Father"}], etag=etag)
    assert ok.status_code == 200, ok.text
    stale = _edit(api, admin, batch_id, 2, [{"column": 2, "value": "Other"}], etag=etag)
    assert stale.status_code == 412
    # Same value again: nothing new is stored, the version stays.
    same = _edit(
        api,
        admin,
        batch_id,
        2,
        [{"column": 2, "value": "Synthetic Father"}],
        etag=ok.headers["ETag"],
    )
    assert same.status_code == 200
    assert same.headers["ETag"] == ok.headers["ETag"]


def test_FR_IMP_008_null_or_blank_clears_a_cell(world: Any, api: Any, admin_engine: Engine) -> None:
    admin = world.person("office_admin")
    batch_id, _ = _batch_with_problem(admin_engine, world)
    etag = _sheet(api, admin, batch_id).headers["ETag"]
    res = _edit(api, admin, batch_id, 2, [{"column": 2, "value": None}], etag=etag)
    assert res.status_code == 200, res.text
    assert res.json()["row"]["cells"][2] == {
        "value": None,
        "edited": True,
        "restricted": False,
        "formula": False,
    }
    res = _edit(api, admin, batch_id, 3, [{"column": 2, "value": "   "}], etag=res.headers["ETag"])
    assert res.status_code == 200, res.text
    assert res.json()["row"]["cells"][2]["value"] is None
    assert _sheet(api, admin, batch_id).json()["edited_cells"] == 2


def test_FR_IMP_008_invariant_6_no_edits_after_the_rows_were_added(
    world: Any, api: Any, admin_engine: Engine
) -> None:
    admin = world.person("office_admin")
    batch_id = S.imported(admin_engine, world.a, S.xlsx_bytes(S.class_list(1)[0]))
    sheet = _sheet(api, admin, batch_id).json()
    assert sheet["editable"] is False
    assert sheet["read_only_reason"] == "committed"
    assert all(not c["editable"] for c in sheet["columns"])
    res = _edit(
        api, admin, batch_id, 2, [{"column": 1, "value": "Changed"}], etag=f'W/"{sheet["version"]}"'
    )
    assert res.status_code == 409
    assert res.json()["code"] == "import_not_editable"
    assert "change request" in res.json()["detail"]
    reverted = api.call(admin, "POST", f"/api/v1/imports/{batch_id}/revert")
    assert reverted.status_code == 200
    again = _sheet(api, admin, batch_id).json()
    assert again["read_only_reason"] == "reverted"
    res = _edit(
        api, admin, batch_id, 2, [{"column": 1, "value": "Changed"}], etag=f'W/"{again["version"]}"'
    )
    assert res.status_code == 409


def test_FR_IMP_008_aadhaar_in_the_file_is_masked_in_the_sheet(
    world: Any, api: Any, admin_engine: Engine
) -> None:
    admin = world.person("office_admin")
    aadhaar = S.valid_aadhaar("34567890123")
    rows, _ = S.class_list(1)
    rows[0].append("Remarks")
    rows[1].append(f"Card {aadhaar}")
    batch_id = S.start(admin_engine, world.a, S.xlsx_bytes(rows))
    res = _sheet(api, admin, batch_id)
    assert aadhaar not in res.text
    remarks = res.json()["data"][0]["cells"][7]["value"]
    assert remarks.startswith("Card ")
    assert aadhaar[-4:] in remarks


def test_FR_IMP_009_export_csv_and_xlsx(world: Any, api: Any, admin_engine: Engine) -> None:
    admin = world.person("office_admin")
    aadhaar = S.valid_aadhaar("45678901234")
    rows, numbers = S.class_list(1)
    rows[0] += ["Remarks", "Religion"]
    rows[1] += [f'=HYPERLINK("x") {aadhaar}', "Synthetic faith"]
    batch_id = S.start(admin_engine, world.a, S.xlsx_bytes(rows))
    etag = _sheet(api, admin, batch_id).headers["ETag"]
    assert (
        _edit(api, admin, batch_id, 2, [{"column": 1, "value": TELUGU}], etag=etag).status_code
        == 200
    )

    res = api.call(admin, "GET", f"/api/v1/imports/{batch_id}/sheet/export")
    assert res.status_code == 200, res.text
    assert res.headers["content-type"].startswith("text/csv")
    assert res.headers["content-disposition"].startswith("attachment;")
    assert res.headers["cache-control"] == "no-store"
    assert res.content.startswith(b"\xef\xbb\xbf")  # UTF-8 BOM: Excel opens Telugu correctly
    text_body = res.content.decode("utf-8-sig")
    assert aadhaar not in text_body
    parsed = list(csv.reader(io.StringIO(text_body)))
    assert parsed[0][:2] == ["Adm No", "Name of the Student"]
    assert parsed[1][0] == numbers[0]
    assert parsed[1][1] == TELUGU  # the edit, round-tripped
    assert parsed[1][7].startswith("'=HYPERLINK")  # formula neutralised (SEC-017)
    # Office admins hold student.read_sensitive: the restricted column is included.
    assert parsed[1][8] == "Synthetic faith"

    xlsx = api.call(admin, "GET", f"/api/v1/imports/{batch_id}/sheet/export?format=xlsx")
    assert xlsx.status_code == 200, xlsx.text
    ws = load_workbook(io.BytesIO(xlsx.content)).active
    assert ws is not None
    values = [[c.value for c in row] for row in ws.iter_rows()]
    assert values[1][1] == TELUGU
    formula_cell = values[1][7]
    assert isinstance(formula_cell, str)
    assert formula_cell.startswith("'=HYPERLINK")
    assert ws.cell(row=2, column=8).data_type == "s"
    assert aadhaar not in repr(values)

    events = _audit(admin_engine, batch_id, "import.sheet_exported")
    assert [e["summary"]["format"] for e in events] == ["csv", "xlsx"]
    assert events[0]["summary"]["restricted_included"] is True
    assert events[0]["summary"]["edited_cells"] == 1

    # Office staff lack student.read_sensitive: restricted cells stay empty.
    staff = world.person("office_staff")
    res = api.call(staff, "GET", f"/api/v1/imports/{batch_id}/sheet/export")
    assert res.status_code == 200, res.text
    assert "Synthetic faith" not in res.content.decode("utf-8-sig")
    assert (
        _audit(admin_engine, batch_id, "import.sheet_exported")[-1]["summary"][
            "restricted_included"
        ]
        is False
    )


def test_FR_EXP_004_export_needs_a_recent_sign_in(
    world: Any, api: Any, admin_engine: Engine
) -> None:
    admin = world.person("office_admin")
    batch_id = S.start(admin_engine, world.a, S.xlsx_bytes(S.class_list(1)[0]))
    stale = api.call(admin, "GET", f"/api/v1/imports/{batch_id}/sheet/export", auth_age_s=301)
    assert stale.status_code == 428
    assert stale.json()["code"] == "step_up_required"
    assert _audit(admin_engine, batch_id, "import.sheet_exported") == []


def test_SEC_003_sheet_needs_import_run_and_stays_in_school(
    world: Any, api: Any, admin_engine: Engine
) -> None:
    batch_id = S.start(admin_engine, world.a, S.xlsx_bytes(S.class_list(1)[0]))
    for role in ("teacher", "accountant", "owner"):
        res = api.call(world.person(role), "GET", f"/api/v1/imports/{batch_id}/sheet")
        assert res.status_code == 403, role
    # School B's owner cannot even name school A's import (and holds no import.run in A).
    b_owner = world.b.people["owner"]
    assert api.call(b_owner, "GET", f"/api/v1/imports/{batch_id}/sheet").status_code in (403, 404)


def test_SEC_012_cell_edits_are_reencrypted_on_key_rotation(
    world: Any, admin_engine: Engine
) -> None:
    from app.core.db import tenant_session
    from app.imports import cells, service
    from app.imports.schemas import RowEditIn
    from app.students import crypto

    batch_id = S.start(admin_engine, world.a, S.xlsx_bytes(S.class_list(1)[0]))
    person = world.a.people["office_admin"]
    with tenant_session(world.a.tenant_id, person.user_id) as s:
        version = S.batch(admin_engine, batch_id)["version"]
        service.edit_row(
            s,
            S.ctx(world.a),
            batch_id,
            2,
            RowEditIn.model_validate({"cells": [{"column": 2, "value": "Synthetic Rekey"}]}),
            expected_version=version,
        )
    with admin_engine.begin() as c:  # pretend the edit was written under an older key
        c.execute(
            text("UPDATE sis.import_cell_edits SET key_version = 99 WHERE batch_id = :b"),
            {"b": batch_id},
        )
    with tenant_session(world.a.tenant_id) as s:
        ring = crypto.get_keyring()
        active = ring.active_version(s)
        assert cells.reencrypt_batch(s, ring, active, 500) >= 1
        assert cells.current_values(s, batch_id)[(2, 2)] == "Synthetic Rekey"
    assert (
        S.count(
            admin_engine,
            "SELECT count(*) FROM sis.import_cell_edits WHERE batch_id = :b AND key_version = 99",
            b=batch_id,
        )
        == 0
    )


def test_CLAUDE_6_5_no_cell_values_in_logs(
    world: Any,
    api: Any,
    admin_engine: Engine,
    capsys: pytest.CaptureFixture[str],
    caplog: pytest.LogCaptureFixture,
) -> None:
    admin = world.person("office_admin")
    batch_id = S.start(admin_engine, world.a, S.xlsx_bytes(S.class_list(1)[0]))
    capsys.readouterr()
    etag = _sheet(api, admin, batch_id).headers["ETag"]
    secret = "Synthetica Logcheck Name"
    aadhaar = S.valid_aadhaar("56789012345")
    assert (
        _edit(api, admin, batch_id, 2, [{"column": 1, "value": secret}], etag=etag).status_code
        == 200
    )
    _edit(api, admin, batch_id, 2, [{"column": 2, "value": aadhaar}], etag=etag)
    api.call(admin, "GET", f"/api/v1/imports/{batch_id}/sheet/export")
    captured = capsys.readouterr()
    logs = captured.out + captured.err + caplog.text
    assert secret not in logs
    assert aadhaar not in logs


def test_CLAUDE_6_4_old_value_of_an_edited_cell_is_stored_masked(
    world: Any, api: Any, admin_engine: Engine
) -> None:
    """The value before an edit comes from the uploaded file: an Aadhaar number in it is
    masked before the history stores it (invariant 4: never stored, not even encrypted)."""
    from app.core.db import tenant_session
    from app.core.redaction import contains_full_aadhaar
    from app.imports import cells

    admin = world.person("office_admin")
    aadhaar = S.valid_aadhaar("67890123456")
    rows, _ = S.class_list(1)
    rows[0].append("Remarks")
    rows[1].append(f"Card {aadhaar}")
    batch_id = S.start(admin_engine, world.a, S.xlsx_bytes(rows))
    etag = _sheet(api, admin, batch_id).headers["ETag"]
    res = _edit(api, admin, batch_id, 2, [{"column": 7, "value": "Card on file"}], etag=etag)
    assert res.status_code == 200, res.text
    with admin_engine.connect() as c:
        edit = c.execute(
            text("SELECT id, old_value_ciphertext FROM sis.import_cell_edits WHERE batch_id = :b"),
            {"b": batch_id},
        ).one()
    with tenant_session(world.a.tenant_id) as s:
        old = cells.decrypt(s, edit.old_value_ciphertext, column=cells.OLD_COLUMN, edit_id=edit.id)
    assert old is not None
    assert old.startswith("Card ")
    assert old.endswith(aadhaar[-4:])
    assert not contains_full_aadhaar(old)


def test_FR_IMP_008_columns_suggested_as_restricted_stay_hidden_when_not_imported(
    world: Any, api: Any, admin_engine: Engine
) -> None:
    """A column whose header names a restricted (C3) field stays hidden even when the office
    chose not to import it (docs/05 §5.2: C3 values are never shown in the staged sheet)."""
    admin = world.person("office_admin")
    rows, _ = S.class_list(1)
    rows[0].append("Religion")
    rows[1].append("Synthetic faith")
    batch_id = S.start(admin_engine, world.a, S.xlsx_bytes(rows))
    batch = S.batch(admin_engine, batch_id)
    columns = [
        {"index": int(k), "target": v} for k, v in batch["mapping"].items() if v != "religion"
    ]
    res = api.call(
        admin,
        "PUT",
        f"/api/v1/imports/{batch_id}/mapping",
        json={"columns": columns},
        headers={"If-Match": f'W/"{batch["version"]}"'},
    )
    assert res.status_code == 200, res.text
    sheet = _sheet(api, admin, batch_id)
    religion = sheet.json()["columns"][7]
    assert religion["target"] is None
    assert religion["restricted"] is True
    assert religion["editable"] is False
    assert "Synthetic faith" not in sheet.text
    refused = _edit(
        api, admin, batch_id, 2, [{"column": 7, "value": "Other"}], etag=sheet.headers["ETag"]
    )
    assert refused.status_code == 422
    assert refused.json()["errors"][0]["code"] == "column_restricted"
    # Staff without student.read_sensitive download it without the column's values.
    staff = world.person("office_staff")
    export = api.call(staff, "GET", f"/api/v1/imports/{batch_id}/sheet/export")
    assert export.status_code == 200, export.text
    assert "Synthetic faith" not in export.content.decode("utf-8-sig")


def test_FR_IMP_007_raw_file_purge_erases_staged_edit_values(
    world: Any, api: Any, admin_engine: Engine
) -> None:
    """Edits hold values from the uploaded file: when the raw file is deleted after the
    retention period, their ciphertext goes too (who edited which cell, and when, stays)."""
    import datetime as dt

    from app.imports import service

    admin = world.person("office_admin")
    batch_id = S.start(admin_engine, world.a, S.xlsx_bytes(S.class_list(1)[0]))
    etag = _sheet(api, admin, batch_id).headers["ETag"]
    res = _edit(api, admin, batch_id, 2, [{"column": 2, "value": "Synthetic Purge"}], etag=etag)
    assert res.status_code == 200, res.text
    assert S.commit(world.a, batch_id) == "committed"
    S.age_batch(
        admin_engine,
        batch_id,
        committed_at=dt.timedelta(days=91),
        revert_deadline=dt.timedelta(days=91),
        created_at=dt.timedelta(days=91),
    )
    assert service.purge_raw_files(world.a.tenant_id) >= 1
    with admin_engine.connect() as c:
        edit = c.execute(
            text(
                "SELECT old_value_ciphertext, new_value_ciphertext, key_version, edited_by, "
                "row_no, column_index FROM sis.import_cell_edits WHERE batch_id = :b"
            ),
            {"b": batch_id},
        ).one()
    assert (edit.old_value_ciphertext, edit.new_value_ciphertext, edit.key_version) == (
        None,
        None,
        None,
    )
    assert (edit.row_no, edit.column_index, edit.edited_by) == (2, 2, admin.user_id)
