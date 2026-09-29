"""Sheets of XLSX/CSV documents: view, save edits as a new version, download (FR-DOC-009,
FR-DOC-010, FR-DOC-011; SEC-017, FR-EXP-003, FR-EXP-004; invariants 4, 5 and 7). Synthetic data
only; Aadhaar-like numbers are built here.

Authorization for every role is in tests/security/test_authz_matrix.py and other-school ids in
tests/security/test_bola.py; these tests cover the behaviour.
"""

from __future__ import annotations

import csv
import io
import sys
import uuid
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import Engine, text

from app.core.redaction import verhoeff_valid
from app.core.spreadsheet import CSV_MIME, XLSX_MIME

pytestmark = pytest.mark.db
S = sys.modules["sos_test_documents_support"]
W = sys.modules["sos_test_api_world"]
DOCS = "/api/v1/documents"
TELUGU = "సింథెటిక విద్యార్థి"


def _aadhaar(body: str = "23456789012") -> str:
    """A synthetic 12-digit number that passes the Verhoeff check (never a real card)."""
    for last in "0123456789":
        if verhoeff_valid(body + last):
            return body + last
    raise AssertionError("no check digit")


def _xlsx(rows: list[list[Any]], *, sheets: int = 1) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Fees"
    for row in rows:
        ws.append(row)
    for i in range(1, sheets):
        wb.create_sheet(f"Other {i}").append(["Synthetic", "Other", "Sheet"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _table() -> list[list[Any]]:
    return [
        ["Receipt", "Name", "Amount", "Remarks"],
        ["R-001", "Synthetica One", 1200, None],
        ["R-002", "Synthetica Two", 950.5, "Paid in cash"],
        ["R-003", "Synthetica Three", 1200, None],
    ]


def _doc(
    world: Any,
    admin: Engine,
    data: bytes,
    *,
    mime: str = XLSX_MIME,
    ext: str = "xlsx",
    sensitivity: str = "C1",
    purpose: str = "circular",
    doc_type: str = "circular",
) -> uuid.UUID:
    doc_id: uuid.UUID = S.make_document(
        admin,
        world.a.tenant_id,
        world.person("owner").user_id,
        data=data,
        mime_type=mime,
        ext=ext,
        sensitivity=sensitivity,
        purpose=purpose,
        doc_type=doc_type,
    )
    return doc_id


def _sheet(api: Any, who: Any, doc_id: uuid.UUID, **params: Any) -> Any:
    return api.call(who, "GET", f"{DOCS}/{doc_id}/sheet", params=params)


def _save(
    api: Any, who: Any, doc_id: uuid.UUID, edits: list[dict[str, Any]], *, etag: str, base: int = 1
) -> Any:
    return api.call(
        who,
        "POST",
        f"{DOCS}/{doc_id}/sheet/versions",
        json={"base_version_no": base, "edits": edits},
        headers={"If-Match": etag},
    )


def _export(api: Any, who: Any, doc_id: uuid.UUID, body: dict[str, Any], **kw: Any) -> Any:
    return api.call(who, "POST", f"{DOCS}/{doc_id}/sheet/export", json=body, **kw)


def _events(admin: Engine, world: Any, doc_id: uuid.UUID, action: str) -> list[dict[str, Any]]:
    return [
        e for e in W.audit_events(admin, world.a.tenant_id, action) if e["resource_id"] == doc_id
    ]


def _versions(admin: Engine, doc_id: uuid.UUID) -> list[Any]:
    with admin.connect() as c:
        return list(
            c.execute(
                text(
                    "SELECT version_no, object_key, mime_type, status FROM kb.document_versions "
                    "WHERE document_id = :d ORDER BY version_no"
                ),
                {"d": doc_id},
            )
        )


# --- view (FR-DOC-009) ---------------------------------------------------------------------------


def test_FR_DOC_009_sheet_shows_the_first_worksheet_masked(
    world: Any, api: Any, admin_engine: Engine
) -> None:
    admin = world.person("office_admin")
    aadhaar = _aadhaar()
    rows = _table()
    rows[1][3] = f"Card {aadhaar}"
    rows[3][3] = '=HYPERLINK("http://example.invalid")'
    doc = _doc(world, admin_engine, _xlsx(rows))
    res = _sheet(api, admin, doc)
    assert res.status_code == 200, res.text
    assert res.headers["cache-control"] == "no-store"
    body = res.json()
    assert res.headers["ETag"] == f'W/"{body["version"]}"'
    assert body["version_no"] == 1
    assert body["kind"] == "xlsx"
    assert body["sheet_count"] == 1
    assert body["total_rows"] == 3
    assert [c["header"] for c in body["columns"]] == ["Receipt", "Name", "Amount", "Remarks"]
    assert [c["letter"] for c in body["columns"]] == ["A", "B", "C", "D"]
    first, second, third = body["data"]
    assert first["row_no"] == 2
    assert first["cells"][2] == {"value": "1200", "formula": False}
    assert second["cells"][2]["value"] == "950.5"
    assert aadhaar not in res.text  # invariant 4
    assert first["cells"][3]["value"].startswith("Card ")
    assert first["cells"][3]["value"].endswith(aadhaar[-4:])
    # Formulas are shown as inert text and never evaluated (SEC-017); such a workbook is
    # read-only here so its formulas are never lost.
    assert third["cells"][3]["formula"] is True
    assert body["editable"] is False
    assert body["read_only_reason"] == "formulas"
    page = _sheet(api, admin, doc, limit=2).json()
    assert [r["row_no"] for r in page["data"]] == [2, 3]
    rest = _sheet(api, admin, doc, limit=2, cursor=page["next_cursor"]).json()
    assert [r["row_no"] for r in rest["data"]] == [4]
    assert rest["next_cursor"] is None


def test_FR_DOC_009_editable_only_for_single_sheet_xlsx_and_uploaders(
    world: Any, api: Any, admin_engine: Engine
) -> None:
    plain = _doc(world, admin_engine, _xlsx(_table()))
    body = _sheet(api, world.person("office_admin"), plain).json()
    assert (body["editable"], body["read_only_reason"]) == (True, None)
    # Readers without document.upload see the sheet read-only.
    reader = _sheet(api, world.person("auditor_readonly"), plain).json()
    assert (reader["editable"], reader["read_only_reason"]) == (False, "no_permission")
    several = _doc(world, admin_engine, _xlsx(_table(), sheets=2))
    body = _sheet(api, world.person("office_admin"), several).json()
    assert (body["sheet_count"], body["read_only_reason"]) == (2, "several_sheets")
    as_csv = _doc(world, admin_engine, S.csv_text(), mime=CSV_MIME, ext="csv")
    body = _sheet(api, world.person("office_admin"), as_csv).json()
    assert (body["kind"], body["read_only_reason"]) == ("csv", "not_versionable")
    assert body["columns"][0]["header"] == "admission_no"


def test_FR_DOC_009_other_files_and_restricted_files(
    world: Any, api: Any, admin_engine: Engine
) -> None:
    admin = world.person("office_admin")
    pdf = _doc(world, admin_engine, S.pdf(), mime="application/pdf", ext="pdf")
    res = _sheet(api, admin, pdf)
    assert res.status_code == 415
    assert res.json()["code"] == "not_a_sheet"
    # C3 files open only for staff who may see sensitive data (as downloads, FR-DOC-004).
    restricted = _doc(world, admin_engine, _xlsx(_table()), sensitivity="C3")
    res = _sheet(api, world.person("office_staff"), restricted)
    assert res.status_code == 403
    assert res.json()["code"] == "sensitive_document"
    assert _sheet(api, admin, restricted).status_code == 200


def test_FR_DOC_009_import_files_open_from_their_import(
    world: Any, api: Any, admin_engine: Engine
) -> None:
    """An import file may hold restricted (C3) columns that only the import sheet hides
    (FR-IMP-008): the document viewer refuses it and points there."""
    doc = _doc(
        world,
        admin_engine,
        _xlsx(_table()),
        sensitivity="C2",
        purpose="import_file",
        doc_type="import_file",
    )
    res = _sheet(api, world.person("office_admin"), doc)
    assert res.status_code == 409
    assert res.json()["code"] == "import_file_sheet"
    res = _export(api, world.person("office_admin"), doc, {"format": "csv"})
    assert res.status_code == 409
    assert res.json()["code"] == "import_file_sheet"


# --- save as a new version (FR-DOC-010) ----------------------------------------------------------


def test_FR_DOC_010_edits_become_the_next_version(
    world: Any, api: Any, admin_engine: Engine
) -> None:
    admin = world.person("office_admin")
    doc = _doc(world, admin_engine, _xlsx(_table()))
    etag = _sheet(api, admin, doc).headers["ETag"]
    res = _save(
        api,
        admin,
        doc,
        [
            {"row_no": 3, "column": 1, "value": f"  {TELUGU}  "},
            {"row_no": 4, "column": 3, "value": "Waived"},
            {"row_no": 4, "column": 2, "value": "1300"},
            {"row_no": 4, "column": 0, "value": "007"},
            {"row_no": 2, "column": 2, "value": None},
        ],
        etag=etag,
    )
    assert res.status_code == 202, res.text
    assert res.headers["Location"] == f"{DOCS}/{doc}"
    versions = _versions(admin_engine, doc)
    assert [(v.version_no, v.mime_type, v.status) for v in versions] == [
        (1, XLSX_MIME, "ready"),
        (2, XLSX_MIME, "queued"),
    ]
    # The stored v1 is untouched; v2 has the edits, unedited numbers keep their type.
    ws = load_workbook(io.BytesIO(S.memory_store().objects[versions[1].object_key].data)).active
    assert ws is not None
    values = [[c.value for c in row] for row in ws.iter_rows()]
    assert values[0] == ["Receipt", "Name", "Amount", "Remarks"]
    assert values[1][2] is None
    assert values[2][1] == TELUGU  # NFC, trimmed
    assert values[2][2] == 950.5
    assert values[3][3] == "Waived"
    assert values[3][2] == 1300  # a number typed into a number cell stays a number
    assert values[3][0] == "007"  # a code stays text (leading zeros kept)
    # The new version is scanned and indexed like any upload (FR-DOC-002).
    registered = S.outbox_events(admin_engine, world.a.tenant_id, "document.version.registered")
    assert len([e for e in registered if e["document_id"] == str(doc)]) == 1
    # Audit: cell references and counts only, never values (invariant 7).
    edited = _events(admin_engine, world, doc, "document.sheet_edited")
    assert len(edited) == 1
    assert edited[0]["summary"]["edited_cells"] == 5
    assert edited[0]["summary"]["cells"] == ["C2", "B3", "A4", "C4", "D4"]
    assert TELUGU not in repr(edited)
    added = _events(admin_engine, world, doc, "document.version_added")
    assert added[-1]["summary"]["source"] == "sheet_editor"
    # While v2 is being checked the sheet shows v1, read-only.
    body = _sheet(api, admin, doc).json()
    assert (body["version_no"], body["read_only_reason"]) == (1, "newer_version")


def test_FR_DOC_010_saving_is_refused_safely(world: Any, api: Any, admin_engine: Engine) -> None:
    admin = world.person("office_admin")
    doc = _doc(world, admin_engine, _xlsx(_table()))
    etag = _sheet(api, admin, doc).headers["ETag"]
    aadhaar = _aadhaar("34567890123")
    res = _save(api, admin, doc, [{"row_no": 2, "column": 3, "value": aadhaar}], etag=etag)
    assert res.status_code == 422
    assert res.json()["errors"][0]["code"] == "aadhaar_full_number_rejected"
    assert aadhaar not in res.text
    bad = _save(api, admin, doc, [{"row_no": 2, "column": 3, "value": "a\nb"}], etag=etag)
    assert bad.json()["errors"][0]["code"] == "control_characters"
    header = _save(api, admin, doc, [{"row_no": 1, "column": 0, "value": "x"}], etag=etag)
    assert header.status_code == 422  # the header row is not edited here
    unknown = _save(api, admin, doc, [{"row_no": 99, "column": 0, "value": "x"}], etag=etag)
    assert unknown.json()["errors"][0]["code"] == "unknown_row"
    stale = _save(api, admin, doc, [{"row_no": 2, "column": 3, "value": "x"}], etag='W/"999"')
    assert stale.status_code == 412
    same = _save(api, admin, doc, [{"row_no": 2, "column": 0, "value": "R-001"}], etag=etag)
    assert same.status_code == 409
    assert same.json()["code"] == "version_unchanged"
    old_base = _save(api, admin, doc, [{"row_no": 2, "column": 3, "value": "x"}], etag=etag, base=2)
    assert old_base.status_code == 404  # there is no version 2
    assert [v.version_no for v in _versions(admin_engine, doc)] == [1]
    reader = _save(
        api,
        world.person("auditor_readonly"),
        doc,
        [{"row_no": 2, "column": 3, "value": "x"}],
        etag=etag,
    )
    assert reader.status_code == 403
    several = _doc(world, admin_engine, _xlsx(_table(), sheets=2))
    etag = _sheet(api, admin, several).headers["ETag"]
    res = _save(api, admin, several, [{"row_no": 2, "column": 3, "value": "x"}], etag=etag)
    assert res.status_code == 409
    assert res.json()["code"] == "workbook_has_several_sheets"


# --- download (FR-DOC-011) -----------------------------------------------------------------------


def test_FR_DOC_011_download_with_unsaved_edits(world: Any, api: Any, admin_engine: Engine) -> None:
    admin = world.person("office_admin")
    aadhaar = _aadhaar("45678901234")
    rows = _table()
    rows[2][3] = f'=HYPERLINK("x") {aadhaar}'
    doc = _doc(world, admin_engine, _xlsx(rows))
    edits = [{"row_no": 2, "column": 1, "value": TELUGU}]
    res = _export(api, admin, doc, {"format": "csv", "edits": edits})
    assert res.status_code == 200, res.text
    assert res.headers["content-type"].startswith("text/csv")
    assert res.headers["content-disposition"].startswith("attachment;")
    assert res.headers["cache-control"] == "no-store"
    assert res.content.startswith(b"\xef\xbb\xbf")
    parsed = list(csv.reader(io.StringIO(res.content.decode("utf-8-sig"))))
    assert parsed[0] == ["Receipt", "Name", "Amount", "Remarks"]
    assert parsed[1][1] == TELUGU
    assert parsed[1][2] == "1200"
    assert parsed[2][3].startswith("'=HYPERLINK")  # SEC-017
    assert aadhaar not in res.content.decode("utf-8-sig")  # invariant 4
    xlsx = _export(api, admin, doc, {"format": "xlsx", "edits": edits})
    assert xlsx.status_code == 200, xlsx.text
    ws = load_workbook(io.BytesIO(xlsx.content)).active
    assert ws is not None
    assert ws.cell(row=2, column=2).value == TELUGU
    assert ws.cell(row=3, column=4).data_type == "s"
    assert str(ws.cell(row=3, column=4).value).startswith("'=")
    page_header = ws.oddHeader
    assert page_header is not None
    assert page_header.center.text  # FR-EXP-003 watermark
    events = _events(admin_engine, world, doc, "document.sheet_exported")
    assert [e["summary"]["format"] for e in events] == ["csv", "xlsx"]
    assert events[0]["summary"]["edited_cells"] == 1
    assert TELUGU not in repr(events)
    # Nothing was stored: still one version.
    assert [v.version_no for v in _versions(admin_engine, doc)] == [1]
    bad = _export(
        api, admin, doc, {"format": "csv", "edits": [{"row_no": 2, "column": 1, "value": aadhaar}]}
    )
    assert bad.status_code == 422
    assert aadhaar not in bad.text


def test_FR_EXP_004_personal_data_downloads_need_a_recent_sign_in(
    world: Any, api: Any, admin_engine: Engine
) -> None:
    admin = world.person("office_admin")
    personal = _doc(world, admin_engine, _xlsx(_table()), sensitivity="C2")
    stale = _export(api, admin, personal, {"format": "csv"}, auth_age_s=301)
    assert stale.status_code == 428
    assert stale.json()["code"] == "step_up_required"
    assert _events(admin_engine, world, personal, "document.sheet_exported") == []
    assert _export(api, admin, personal, {"format": "csv"}).status_code == 200
    # Internal (C1) sheets hold no personal data: no step-up.
    internal = _doc(world, admin_engine, _xlsx(_table()), sensitivity="C1")
    assert _export(api, admin, internal, {"format": "csv"}, auth_age_s=301).status_code == 200


def test_CLAUDE_6_5_no_sheet_values_in_logs(
    world: Any,
    api: Any,
    admin_engine: Engine,
    capsys: pytest.CaptureFixture[str],
    caplog: pytest.LogCaptureFixture,
) -> None:
    admin = world.person("office_admin")
    doc = _doc(world, admin_engine, _xlsx(_table()))
    capsys.readouterr()
    secret = "Synthetica Logcheck Payer"
    aadhaar = _aadhaar("56789012345")
    etag = _sheet(api, admin, doc).headers["ETag"]
    _export(
        api, admin, doc, {"format": "xlsx", "edits": [{"row_no": 2, "column": 1, "value": secret}]}
    )
    _save(api, admin, doc, [{"row_no": 2, "column": 3, "value": aadhaar}], etag=etag)
    assert (
        _save(api, admin, doc, [{"row_no": 2, "column": 1, "value": secret}], etag=etag).status_code
        == 202
    )
    captured = capsys.readouterr()
    logs = captured.out + captured.err + caplog.text
    assert secret not in logs
    assert "Synthetica One" not in logs
    assert aadhaar not in logs
