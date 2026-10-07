"""Safe spreadsheet reading and writing shared by imports and documents (SEC-017, invariant 4,
docs/07 §10; FR-IMP-001, FR-IMP-008, FR-IMP-009, FR-DOC-009..011). Pure: no database, no web.

Reading (the rules imports have always used, moved here so documents apply the same ones):

- **XLSX** is opened with openpyxl ``read_only=True, data_only=False`` after a zip pre-check
  (member count, declared uncompressed size, compression ratio: zip bombs are refused before any
  XML is parsed). openpyxl parses XML through defusedxml when it is installed (entity expansion
  and external entities are refused). ``data_only=False`` means a formula cell yields its
  formula text; nothing is ever evaluated or recalculated, and cached results are ignored. Only
  the first visible worksheet is read; :func:`xlsx_sheet_count` says how many there are.
- **CSV** must be UTF-8 (a BOM is accepted); the delimiter is sniffed among ``, ; TAB |``.
- **Formulas** (formula cells, or text starting with ``=``/``@``, or ``+``/``-`` followed by
  something other than a number) are kept as inert literal text and flagged.

Writing (:func:`write_csv`, :func:`write_xlsx`): every value passes :func:`safe_cell`: NFC,
control characters removed, any 12-digit Verhoeff-valid sequence masked
(:func:`app.core.redaction.mask_aadhaar`, invariant 4) and formula injection neutralised with a
leading apostrophe (cells starting with ``=``, ``+``, ``-``, ``@``, tab or carriage return, also
after leading whitespace or as full-width ``＝＋－＠``; OWASP CSV injection). CSV is UTF-8 with a BOM so Excel opens Telugu correctly; XLSX cells are written
as explicit strings (type ``s``), so a value is never stored as a formula.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import re
import unicodedata
import zipfile
from collections.abc import Iterator, Sequence
from dataclasses import dataclass
from typing import Final, Literal, Protocol

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.core.redaction import mask_aadhaar

FileKind = Literal["xlsx", "csv"]
CellValue = str | int | float | bool | dt.date | dt.datetime | None

XLSX_MIME: Final = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV_MIME: Final = "text/csv"
CSV_BOM: Final = "﻿"  # Excel opens UTF-8 (Telugu) correctly only with a BOM
FORMULA_TRIGGERS: Final = ("=", "+", "-", "@", "\t", "\r")
NEUTRALISER: Final = "'"
MAX_XLSX_CELL_CHARS: Final = 32_000  # XLSX allows 32,767 characters per cell
MAX_ROW_CELLS: Final = 16_384  # XLSX's last column is XFD; no real row holds more cells

_NUMERIC_RE: Final = re.compile(r"^[+-]?[\d\s().,-]*$")
_CSV_DELIMITERS: Final = ",;\t|"
_CONTROL_RE: Final = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
_SHEET_NAME_BAD: Final = re.compile(r"[\[\]:*?/\\]")


class SpreadsheetError(Exception):
    """The file cannot be read as a sheet (``code`` is shown to the user)."""

    def __init__(self, code: str) -> None:
        super().__init__(code)
        self.code = code


class ReadLimits(Protocol):
    """The limits a reader needs (``app.imports.config.Limits`` and the documents viewer
    limits both satisfy it)."""

    @property
    def max_scanned_rows(self) -> int: ...
    @property
    def xlsx_max_uncompressed_bytes(self) -> int: ...
    @property
    def xlsx_max_compression_ratio(self) -> int: ...
    @property
    def xlsx_max_members(self) -> int: ...


@dataclass(frozen=True, slots=True)
class Cell:
    value: CellValue
    formula: bool = False

    @property
    def empty(self) -> bool:
        return self.value is None or (isinstance(self.value, str) and not self.value.strip())


def _formula_head(text: str) -> str:
    """``text`` without leading whitespace, control and invisible format characters (spreadsheet
    programs skip them), with its first character NFKC-folded so full-width ``＝＋－＠`` read as
    ``=+-@`` (audit 2026-10-04 data-layer hardening note 4)."""
    i = 0
    while i < len(text) and (text[i].isspace() or unicodedata.category(text[i]) in ("Cc", "Cf")):
        i += 1
    rest = text[i:]
    if not rest:
        return ""
    return unicodedata.normalize("NFKC", rest[0]) + rest[1:]


def starts_formula(text: str) -> bool:
    """True when a spreadsheet program could read ``text`` as a formula: it starts with one of
    :data:`FORMULA_TRIGGERS`, or does so once leading whitespace, control and format characters
    are skipped and the first sign is NFKC-folded (SEC-017, OWASP CSV injection)."""
    return text.startswith(FORMULA_TRIGGERS) or _formula_head(text).startswith(FORMULA_TRIGGERS)


def looks_like_formula(text: str) -> bool:
    """Text a spreadsheet program would treat as a formula or DDE payload."""
    stripped = _formula_head(text)
    if not stripped:
        return False
    if stripped[0] in "=@":
        return True
    return stripped[0] in "+-" and _NUMERIC_RE.match(stripped) is None


def _normalise_text(value: str) -> str:
    return unicodedata.normalize("NFC", value)


# --- XLSX ---------------------------------------------------------------------------------------


def check_zip(data: bytes, limits: ReadLimits) -> None:
    """Refuse archives that are not zips or that would inflate too far (zip bombs)."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
    except (zipfile.BadZipFile, ValueError) as exc:
        raise SpreadsheetError("file_unreadable") from exc
    if len(members) > limits.xlsx_max_members:
        raise SpreadsheetError("file_too_complex")
    total = 0
    for info in members:
        total += info.file_size
        if total > limits.xlsx_max_uncompressed_bytes:
            raise SpreadsheetError("file_too_complex")
        ratio = info.file_size / max(info.compress_size, 1)
        if info.file_size > 1_000_000 and ratio > limits.xlsx_max_compression_ratio:
            raise SpreadsheetError("file_too_complex")
    _check_row_widths(data)


_ROW_OR_CELL_RE: Final = re.compile(rb"<(?:[A-Za-z_][\w.-]*:)?(row|c)[\s/>]")
_SCAN_CHUNK: Final = 1 << 20
_SCAN_OVERLAP: Final = 64


def _check_row_widths(data: bytes) -> None:
    """Refuse a worksheet row holding more than :data:`MAX_ROW_CELLS` cell elements.

    openpyxl builds a whole ``<row>`` (every ``<c>`` element) before yielding it, so one row of
    millions of small cells inside the size and ratio limits costs gigabytes. This streams every
    member (a worksheet part may have any name; the workbook relationships choose it) and counts
    ``<c>`` start tags since the last ``<row>`` tag without parsing XML. Text never contains a
    raw ``<``, so only markup is counted (SEC-017)."""
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        for info in archive.infolist():
            if info.is_dir():
                continue
            cells = 0
            pending = b""
            try:
                with archive.open(info) as member:
                    while True:
                        chunk = member.read(_SCAN_CHUNK)
                        window = pending + chunk
                        # A tag starting in the last bytes may be cut: look at it next round.
                        end = len(window) - _SCAN_OVERLAP if chunk else len(window)
                        for match in _ROW_OR_CELL_RE.finditer(window, 0, len(window)):
                            if match.start() >= end:
                                break
                            if match.group(1) == b"row":
                                cells = 0
                            else:
                                cells += 1
                                if cells > MAX_ROW_CELLS:
                                    raise SpreadsheetError("too_many_columns")
                        if not chunk:
                            break
                        pending = window[max(end, 0) :]
            except (zipfile.BadZipFile, NotImplementedError, RuntimeError, OSError) as exc:
                raise SpreadsheetError("file_unreadable") from exc


def xlsx_rows(data: bytes, limits: ReadLimits) -> Iterator[tuple[int, list[Cell]]]:
    """``(row number, cells)`` of the first visible worksheet; formulas never evaluated."""
    check_zip(data, limits)
    try:
        workbook = load_workbook(
            io.BytesIO(data), read_only=True, data_only=False, keep_links=False
        )
    except Exception as exc:  # openpyxl raises many types for damaged/hostile files
        raise SpreadsheetError("file_unreadable") from exc
    try:
        sheets = [
            ws for ws in workbook.worksheets if getattr(ws, "sheet_state", "visible") == "visible"
        ]
        if not sheets:
            raise SpreadsheetError("no_worksheet")
        ws = sheets[0]
        reset = getattr(ws, "reset_dimensions", None)
        if callable(reset):
            reset()  # do not trust the stored dimension; read to the real end
        try:
            for row_no, row in enumerate(ws.iter_rows(), start=1):
                if row_no > limits.max_scanned_rows:
                    raise SpreadsheetError("too_many_rows")
                cells: list[Cell] = []
                for cell in row:
                    value = getattr(cell, "value", None)
                    is_formula = getattr(cell, "data_type", None) == "f"
                    if isinstance(value, str):
                        value = _normalise_text(value)
                        is_formula = is_formula or looks_like_formula(value)
                    elif isinstance(value, dt.time | dt.timedelta):
                        value = str(value)
                    elif value is not None and not isinstance(
                        value, int | float | bool | dt.date | dt.datetime
                    ):
                        value = str(value)  # e.g. rich text / array formula objects: inert text
                    cells.append(Cell(value, is_formula))
                yield row_no, cells
        except SpreadsheetError:
            raise
        except Exception as exc:  # damaged XML, forbidden entities (defusedxml), bad values
            raise SpreadsheetError("file_unreadable") from exc
    finally:
        workbook.close()


def xlsx_sheet_count(data: bytes, limits: ReadLimits) -> int:
    """How many worksheets the workbook has (visible or not; charts sheets excluded)."""
    check_zip(data, limits)
    try:
        workbook = load_workbook(
            io.BytesIO(data), read_only=True, data_only=False, keep_links=False
        )
    except Exception as exc:
        raise SpreadsheetError("file_unreadable") from exc
    try:
        return len(workbook.worksheets)
    finally:
        workbook.close()


# --- CSV ----------------------------------------------------------------------------------------


def csv_rows(data: bytes, limits: ReadLimits) -> Iterator[tuple[int, list[Cell]]]:
    """``(row number, cells)`` of a UTF-8 CSV (BOM accepted, delimiter sniffed)."""
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise SpreadsheetError("not_utf8") from exc
    if "\x00" in text:
        raise SpreadsheetError("file_unreadable")
    sample = text[:65536]
    try:
        dialect: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(
            sample, delimiters=_CSV_DELIMITERS
        )
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text, newline=""), dialect)
    try:
        for row_no, raw in enumerate(reader, start=1):
            if row_no > limits.max_scanned_rows:
                raise SpreadsheetError("too_many_rows")
            if len(raw) > MAX_ROW_CELLS:
                raise SpreadsheetError("too_many_columns")
            cells = []
            for value in raw:
                clean = _normalise_text(value)
                cells.append(Cell(clean if clean.strip() else None, looks_like_formula(clean)))
            yield row_no, cells
    except csv.Error as exc:
        raise SpreadsheetError("file_unreadable") from exc


def raw_rows(data: bytes, kind: FileKind, limits: ReadLimits) -> Iterator[tuple[int, list[Cell]]]:
    return xlsx_rows(data, limits) if kind == "xlsx" else csv_rows(data, limits)


# --- display ------------------------------------------------------------------------------------


def display_text(value: CellValue) -> str | None:
    """A cell as display text (``None`` when empty): dates in ISO format, whole floats without
    ``.0``, NFC text. Not masked: callers mask with :func:`app.core.redaction.mask_aadhaar`."""
    if value is None:
        return None
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time() else value.isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    text = unicodedata.normalize("NFC", str(value))
    return text if text.strip() else None


# --- writing ------------------------------------------------------------------------------------


def neutralise_formula(text: str) -> str:
    """Prefix ``'`` when ``text`` would start a formula (SEC-017, OWASP CSV injection), also
    behind leading whitespace or a full-width sign (:func:`starts_formula`)."""
    return NEUTRALISER + text if starts_formula(text) else text


def safe_cell(value: object) -> str:
    """The text written to an exported cell: NFC, no control characters (tab, LF and CR are
    kept so the formula check sees them), Aadhaar-like numbers masked, formula neutralised."""
    if value is None:
        return ""
    text = unicodedata.normalize("NFC", str(value))
    text = _CONTROL_RE.sub("", text)
    text = mask_aadhaar(text)[:MAX_XLSX_CELL_CHARS]
    return neutralise_formula(text)


def write_csv(header: Sequence[object], rows: Sequence[Sequence[object]]) -> bytes:
    """UTF-8 CSV with a BOM: the header row, then the rows, every cell through
    :func:`safe_cell`."""
    out = io.StringIO()
    out.write(CSV_BOM)
    writer = csv.writer(out, lineterminator="\r\n", quoting=csv.QUOTE_MINIMAL)
    writer.writerow([safe_cell(h) for h in header])
    for row in rows:
        writer.writerow([safe_cell(v) for v in row])
    return out.getvalue().encode("utf-8")


def sheet_title(name: str) -> str:
    """A valid worksheet title (at most 31 characters, no ``[]:*?/\\``)."""
    return _SHEET_NAME_BAD.sub(" ", unicodedata.normalize("NFC", name)).strip()[:31] or "Sheet"


_HEADER_FONT: Final = Font(bold=True)
_HEADER_FILL: Final = PatternFill("solid", fgColor="F1F1F1")


def _put(ws: object, row: int, col: int, value: object) -> None:
    cell = ws.cell(row=row, column=col)  # type: ignore[attr-defined]
    cell.value = safe_cell(value)
    cell.data_type = "s"  # never a formula, whatever the text
    cell.number_format = "@"


def write_xlsx(
    header: Sequence[object],
    rows: Sequence[Sequence[object]],
    *,
    title: str,
    watermark: str | None = None,
) -> bytes:
    """One worksheet: bold frozen header row, text cells, optional watermark in the printed page
    header and the file properties."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None  # noqa: S101 - a new workbook always has one sheet
    ws.title = sheet_title(title)
    wb.properties.creator = "SchoolOS"
    wb.properties.title = safe_cell(title)
    if watermark:
        wb.properties.description = safe_cell(watermark)
        page_header = ws.oddHeader
        if page_header is not None:
            page_header.center.text = safe_cell(watermark).replace("&", "&&")
    for col, text in enumerate(header, start=1):
        _put(ws, 1, col, text)
        ws.cell(row=1, column=col).font = _HEADER_FONT
        ws.cell(row=1, column=col).fill = _HEADER_FILL
    widths = [len(str(h or "")) for h in header]
    for r, row in enumerate(rows, start=2):
        for col, value in enumerate(row, start=1):
            _put(ws, r, col, value)
            if r <= 500 and col <= len(widths):
                widths[col - 1] = max(widths[col - 1], len(str(value or "")))
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 60)
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def column_letter(index: int) -> str:
    """Spreadsheet column name of a 0-based index (0 -> ``A``, 26 -> ``AA``)."""
    return get_column_letter(index + 1)


__all__ = [
    "CSV_BOM",
    "CSV_MIME",
    "FORMULA_TRIGGERS",
    "XLSX_MIME",
    "Cell",
    "CellValue",
    "FileKind",
    "ReadLimits",
    "SpreadsheetError",
    "check_zip",
    "column_letter",
    "csv_rows",
    "display_text",
    "looks_like_formula",
    "neutralise_formula",
    "raw_rows",
    "safe_cell",
    "sheet_title",
    "write_csv",
    "write_xlsx",
    "xlsx_rows",
    "xlsx_sheet_count",
]
