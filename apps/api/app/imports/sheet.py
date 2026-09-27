"""Read an uploaded spreadsheet into plain cells (FR-IMP-001, SEC-017, docs/07 §10). Pure.

- **XLSX** is opened with openpyxl ``read_only=True, data_only=False`` after a zip pre-check
  (member count, declared uncompressed size, compression ratio: zip bombs are refused before any
  XML is parsed). openpyxl parses XML through defusedxml when it is installed (entity expansion
  and external entities are refused). ``data_only=False`` means a formula cell yields its
  formula text; nothing is ever evaluated or recalculated, and cached results are ignored.
- **CSV** (and Google Sheets CSV exports) must be UTF-8 (a BOM is accepted); the delimiter is
  sniffed among ``, ; TAB |``.
- **Formulas** (formula cells, or text starting with ``=``/``@``, or ``+``/``-`` followed by
  something other than a number) are kept as inert literal text and flagged
  (``formula_not_evaluated``); they never become values.
- The header is the first of the first rows with at least three text cells. Limits: rows,
  columns, cell length and physical rows scanned (config.yaml).
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import re
import unicodedata
import zipfile
from collections.abc import Iterable, Iterator, Sequence
from dataclasses import dataclass
from typing import Final, Literal

from openpyxl import load_workbook

from app.imports.config import Limits

FileKind = Literal["xlsx", "csv"]
CellValue = str | int | float | bool | dt.date | dt.datetime | None

_NUMERIC_RE: Final = re.compile(r"^[+-]?[\d\s().,-]*$")
_CSV_DELIMITERS: Final = ",;\t|"


class SheetError(Exception):
    """The file cannot be imported as a whole (``code`` is shown to the user)."""

    def __init__(self, code: str) -> None:
        super().__init__(code)
        self.code = code


@dataclass(frozen=True, slots=True)
class Cell:
    value: CellValue
    formula: bool = False

    @property
    def empty(self) -> bool:
        return self.value is None or (isinstance(self.value, str) and not self.value.strip())


@dataclass(frozen=True, slots=True)
class SheetRow:
    row_no: int  # 1-based row number as the spreadsheet shows it
    cells: tuple[Cell, ...]

    def cell(self, index: int) -> Cell:
        return self.cells[index] if index < len(self.cells) else Cell(None)


@dataclass(frozen=True, slots=True)
class Sheet:
    kind: FileKind
    header_row: int
    headers: tuple[str, ...]
    rows: tuple[SheetRow, ...]
    formula_cells: int


def looks_like_formula(text: str) -> bool:
    """Text a spreadsheet program would treat as a formula or DDE payload."""
    stripped = text.lstrip()
    if not stripped:
        return False
    if stripped[0] in "=@":
        return True
    return stripped[0] in "+-" and _NUMERIC_RE.match(stripped) is None


def _normalise_text(value: str) -> str:
    return unicodedata.normalize("NFC", value)


# --- XLSX ---------------------------------------------------------------------------------------


def _check_zip(data: bytes, limits: Limits) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
    except (zipfile.BadZipFile, ValueError) as exc:
        raise SheetError("file_unreadable") from exc
    if len(members) > limits.xlsx_max_members:
        raise SheetError("file_too_complex")
    total = 0
    for info in members:
        total += info.file_size
        if total > limits.xlsx_max_uncompressed_bytes:
            raise SheetError("file_too_complex")
        ratio = info.file_size / max(info.compress_size, 1)
        if info.file_size > 1_000_000 and ratio > limits.xlsx_max_compression_ratio:
            raise SheetError("file_too_complex")


def _xlsx_rows(data: bytes, limits: Limits) -> Iterator[tuple[int, list[Cell]]]:
    _check_zip(data, limits)
    try:
        workbook = load_workbook(
            io.BytesIO(data), read_only=True, data_only=False, keep_links=False
        )
    except Exception as exc:  # openpyxl raises many types for damaged/hostile files
        raise SheetError("file_unreadable") from exc
    try:
        sheets = [
            ws for ws in workbook.worksheets if getattr(ws, "sheet_state", "visible") == "visible"
        ]
        if not sheets:
            raise SheetError("no_worksheet")
        ws = sheets[0]
        reset = getattr(ws, "reset_dimensions", None)
        if callable(reset):
            reset()  # do not trust the stored dimension; read to the real end
        try:
            for row_no, row in enumerate(ws.iter_rows(), start=1):
                if row_no > limits.max_scanned_rows:
                    raise SheetError("too_many_rows")
                cells: list[Cell] = []
                for cell in row:
                    value = getattr(cell, "value", None)
                    is_formula = getattr(cell, "data_type", None) == "f"
                    if isinstance(value, str):
                        value = _normalise_text(value)
                        is_formula = is_formula or looks_like_formula(value)
                    elif isinstance(value, dt.time | dt.timedelta):
                        value = str(value)
                    elif value is not None and not isinstance(
                        value, int | float | bool | dt.date | dt.datetime
                    ):
                        value = str(value)  # e.g. rich text / array formula objects: inert text
                    cells.append(Cell(value, is_formula))
                yield row_no, cells
        except SheetError:
            raise
        except Exception as exc:  # damaged XML, forbidden entities (defusedxml), bad values
            raise SheetError("file_unreadable") from exc
    finally:
        workbook.close()


# --- CSV ----------------------------------------------------------------------------------------


def _csv_rows(data: bytes, limits: Limits) -> Iterator[tuple[int, list[Cell]]]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise SheetError("not_utf8") from exc
    if "\x00" in text:
        raise SheetError("file_unreadable")
    sample = text[:65536]
    try:
        dialect: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(
            sample, delimiters=_CSV_DELIMITERS
        )
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text, newline=""), dialect)
    try:
        for row_no, raw in enumerate(reader, start=1):
            if row_no > limits.max_scanned_rows:
                raise SheetError("too_many_rows")
            cells = []
            for value in raw:
                clean = _normalise_text(value)
                cells.append(Cell(clean if clean.strip() else None, looks_like_formula(clean)))
            yield row_no, cells
    except csv.Error as exc:
        raise SheetError("file_unreadable") from exc


# --- common ---------------------------------------------------------------------------------------


def _is_text(cell: Cell) -> bool:
    return isinstance(cell.value, str) and bool(cell.value.strip()) and not cell.formula


def _trim(cells: Sequence[Cell]) -> list[Cell]:
    end = len(cells)
    while end and cells[end - 1].empty:
        end -= 1
    return list(cells[:end])


def _check_cells(cells: Sequence[Cell], limits: Limits) -> None:
    if len(cells) > limits.max_columns:
        raise SheetError("too_many_columns")
    for cell in cells:
        if isinstance(cell.value, str) and len(cell.value) > limits.max_cell_chars:
            raise SheetError("cell_too_long")


def build_sheet(
    kind: FileKind, raw_rows: Iterable[tuple[int, list[Cell]]], limits: Limits
) -> Sheet:
    header_row = 0
    headers: tuple[str, ...] = ()
    rows: list[SheetRow] = []
    formulas = 0
    for row_no, raw in raw_rows:
        cells = _trim(raw)
        if not header_row:
            if row_no > limits.header_search_rows:
                raise SheetError("header_not_found")
            if sum(1 for c in cells if _is_text(c)) >= limits.header_min_text_cells:
                _check_cells(cells, limits)
                header_row = row_no
                headers = tuple(
                    " ".join(str(c.value).split()) if not c.empty else "" for c in cells
                )
            continue
        if not cells:
            continue
        _check_cells(cells, limits)
        formulas += sum(1 for c in cells if c.formula)
        rows.append(SheetRow(row_no, tuple(cells)))
        if len(rows) > limits.max_rows:
            raise SheetError("too_many_rows")
    if not header_row:
        raise SheetError("header_not_found")
    if not rows:
        raise SheetError("no_data_rows")
    width = max(len(headers), *(len(r.cells) for r in rows))
    if width > limits.max_columns:
        raise SheetError("too_many_columns")
    headers = headers + ("",) * (width - len(headers))
    return Sheet(kind, header_row, headers, tuple(rows), formulas)


def read_sheet(data: bytes, kind: FileKind, limits: Limits) -> Sheet:
    """Parse ``data`` (never evaluating formulas) or raise :class:`SheetError`."""
    if len(data) > limits.max_file_bytes:
        raise SheetError("file_too_large")
    if not data:
        raise SheetError("empty_file")
    rows = _xlsx_rows(data, limits) if kind == "xlsx" else _csv_rows(data, limits)
    return build_sheet(kind, rows, limits)


__all__ = [
    "Cell",
    "CellValue",
    "FileKind",
    "Sheet",
    "SheetError",
    "SheetRow",
    "looks_like_formula",
    "read_sheet",
]
