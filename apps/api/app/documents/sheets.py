"""Sheet grids of XLSX/CSV documents (FR-DOC-009..011; SEC-017, invariant 4). Pure.

- :func:`read_grid` reads the first worksheet of an XLSX (or a CSV) with the shared safe readers
  of :mod:`app.core.spreadsheet` (zip-bomb pre-check, defusedxml, formulas never evaluated) and
  the limits in ``sheets.yaml``. Row 1 is the header row; physical row numbers are kept.
- :func:`apply_edits` puts edited values over the grid (the stored file never changes).
- :func:`build_xlsx` makes the bytes of a new version: one worksheet, values only. Unedited
  cells keep their type (numbers and dates stay numbers and dates); an edited cell keeps the
  type of the cell it replaces when the new text is that type's display form, else it is text;
  every text cell is written as an explicit string (never a formula) and any Aadhaar-like number
  is masked (invariant 4). Workbooks with more sheets or with formulas are never rebuilt (the
  service refuses them), so nothing but formatting is lost, and the UI says so.
"""

from __future__ import annotations

import datetime as dt
import io
import unicodedata
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from functools import lru_cache
from importlib import resources
from typing import Any, Final

import yaml
from openpyxl import Workbook

from app.core.languages import shown_texts
from app.core.redaction import mask_aadhaar
from app.core.spreadsheet import (
    Cell,
    CellValue,
    FileKind,
    SpreadsheetError,
    display_text,
    raw_rows,
    sheet_title,
    xlsx_sheet_count,
)

CellKey = tuple[int, int]  # (row number as the sheet shows it, 0-based column)
_MAX_XLSX_CHARS: Final = 32_000


@dataclass(frozen=True, slots=True)
class SheetLimits:
    max_file_bytes: int
    max_rows: int
    max_columns: int
    max_cell_chars: int
    max_scanned_rows: int
    xlsx_max_uncompressed_bytes: int
    xlsx_max_compression_ratio: int
    xlsx_max_members: int


@dataclass(frozen=True, slots=True)
class SheetConfig:
    limits: SheetLimits
    max_edit_cells: int
    max_value_chars: int
    watermark_texts: Mapping[str, str]

    @property
    def watermark(self) -> str:
        """The download watermark in the languages shown (English only while Telugu is
        hidden, ADR-0036)."""
        return shown_texts(self.watermark_texts)


@lru_cache(maxsize=1)
def sheet_config() -> SheetConfig:
    raw: dict[str, Any] = yaml.safe_load(
        resources.files("app.documents").joinpath("sheets.yaml").read_text("utf-8")
    )
    edits = raw["edits"]
    return SheetConfig(
        limits=SheetLimits(**{k: int(v) for k, v in raw["limits"].items()}),
        max_edit_cells=int(edits["max_cells"]),
        max_value_chars=int(edits["max_value_chars"]),
        watermark_texts={str(k): str(v) for k, v in raw["export_watermark"].items()},
    )


@dataclass(frozen=True, slots=True)
class Grid:
    """Rows of the first worksheet (``rows[0]`` is spreadsheet row 1, the header row)."""

    kind: FileKind
    rows: tuple[tuple[Cell, ...], ...]
    width: int
    sheet_count: int
    formula_cells: int

    def cell(self, row_no: int, column: int) -> Cell:
        if not 1 <= row_no <= len(self.rows):
            return Cell(None)
        row = self.rows[row_no - 1]
        return row[column] if column < len(row) else Cell(None)

    @property
    def data_rows(self) -> int:
        return max(len(self.rows) - 1, 0)


def _trim(cells: Sequence[Cell]) -> tuple[Cell, ...]:
    end = len(cells)
    while end and cells[end - 1].empty:
        end -= 1
    return tuple(cells[:end])


def read_grid(data: bytes, kind: FileKind, limits: SheetLimits) -> Grid:
    """Parse ``data`` (never evaluating formulas) or raise :class:`SpreadsheetError`."""
    if not data:
        raise SpreadsheetError("empty_file")
    if len(data) > limits.max_file_bytes:
        raise SpreadsheetError("file_too_large")
    rows: list[tuple[Cell, ...]] = []
    last_filled = 0
    formulas = 0
    for _row_no, raw in raw_rows(data, kind, limits):
        cells = _trim(raw)
        if len(cells) > limits.max_columns:
            raise SpreadsheetError("too_many_columns")
        for cell in cells:
            if isinstance(cell.value, str) and len(cell.value) > limits.max_cell_chars:
                raise SpreadsheetError("cell_too_long")
        rows.append(cells)
        if cells:
            last_filled = len(rows)
            if last_filled > limits.max_rows:
                raise SpreadsheetError("too_many_rows")
        formulas += sum(1 for c in cells if c.formula)
    rows = rows[:last_filled]
    if not rows:
        raise SpreadsheetError("empty_sheet")
    width = max(len(r) for r in rows)
    count = xlsx_sheet_count(data, limits) if kind == "xlsx" else 1
    return Grid(kind, tuple(rows), width, count, formulas)


def changed_cells(grid: Grid, edits: Mapping[CellKey, str | None]) -> dict[CellKey, str | None]:
    """The edits whose value differs from what the cell shows now (display text)."""
    return {
        key: value for key, value in edits.items() if display_text(grid.cell(*key).value) != value
    }


def _typed(before: CellValue, text: str) -> CellValue:
    """An edited value keeps the type of the cell it replaces when the text is that type's
    display form exactly (``1300`` in a number cell stays a number, ``2026-06-01`` in a date
    cell stays a date); anything else is text, so leading zeros and codes are never lost."""
    candidate: CellValue = None
    try:
        if isinstance(before, bool):
            candidate = None
        elif isinstance(before, int | float):
            number = float(text)
            candidate = int(number) if number.is_integer() and "." not in text else number
        elif isinstance(before, dt.datetime):
            candidate = dt.datetime.combine(dt.date.fromisoformat(text), dt.time())
        elif isinstance(before, dt.date):
            candidate = dt.date.fromisoformat(text)
    except ValueError:
        candidate = None
    if candidate is not None and display_text(candidate) == text:
        return candidate
    return text


def apply_edits(grid: Grid, edits: Mapping[CellKey, str | None]) -> Grid:
    """``grid`` with edited values (``None`` or blank clears; see :func:`_typed` for types).
    Unknown rows are ignored."""
    if not edits:
        return grid
    rows = [list(r) for r in grid.rows]
    width = grid.width
    for (row_no, column), value in edits.items():
        if not 1 <= row_no <= len(rows):
            continue
        row = rows[row_no - 1]
        if column >= len(row):
            row.extend([Cell(None)] * (column + 1 - len(row)))
        if value is None or not value.strip():
            row[column] = Cell(None)
        else:
            row[column] = Cell(_typed(row[column].value, value))
        width = max(width, column + 1)
    cells = tuple(tuple(r) for r in rows)
    return Grid(grid.kind, cells, width, grid.sheet_count, grid.formula_cells)


def cell_display(cell: Cell) -> str | None:
    """Display text with Aadhaar-like numbers masked (invariant 4)."""
    text = display_text(cell.value)
    return mask_aadhaar(text) if text is not None else None


def _stored_value(value: CellValue) -> CellValue:
    """What a rebuilt workbook stores: typed values as they were, text NFC with Aadhaar-like
    numbers masked (a number that masks becomes masked text)."""
    if value is None or isinstance(value, bool | dt.date | dt.datetime):
        return value
    if isinstance(value, int | float):
        text = display_text(value) or ""
        masked = mask_aadhaar(text)
        return value if masked == text else masked
    return mask_aadhaar(unicodedata.normalize("NFC", str(value)))[:_MAX_XLSX_CHARS]


def build_xlsx(grid: Grid, *, title: str) -> bytes:
    """The bytes of a new version: one worksheet with the grid's values (see module doc)."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None  # noqa: S101 - a new workbook always has one sheet
    ws.title = sheet_title(title)
    wb.properties.creator = "SchoolOS"
    for r, row in enumerate(grid.rows, start=1):
        for c, cell in enumerate(row, start=1):
            value = _stored_value(cell.value)
            if value is None:
                continue
            target = ws.cell(row=r, column=c)
            target.value = value
            if isinstance(value, str):
                target.data_type = "s"  # never a formula, whatever the text
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


__all__ = [
    "CellKey",
    "Grid",
    "SheetConfig",
    "SheetLimits",
    "apply_edits",
    "build_xlsx",
    "cell_display",
    "changed_cells",
    "read_grid",
    "sheet_config",
]
